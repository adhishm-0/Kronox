"""Format-aware parsers. Optional packages are imported only for their formats."""
import csv
import html
import json
import logging
import re
import zipfile
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

from .base import Parser, block

logger = logging.getLogger(__name__)


class TextParser(Parser):
    extensions = (".txt", ".md", ".markdown", ".html", ".htm", ".rtf", ".yaml", ".yml", ".json", ".xml")

    def parse(self, path: str, name: str) -> list[dict[str, Any]]:
        ext = Path(name).suffix.lower()
        raw = Path(path).read_text(encoding="utf-8", errors="replace")
        if ext in (".json",):
            obj = json.loads(raw)
            return _json_blocks(obj)
        if ext in (".xml",):
            root = ET.fromstring(raw)
            return _xml_blocks(root)
        if ext in (".html", ".htm"):
            raw = re.sub(r"(?is)<(script|style).*?>.*?</\1>", " ", raw)
            raw = re.sub(r"(?i)</(h[1-6]|p|div|li|tr|br|section)\s*>", "\n", raw)
            text = html.unescape(re.sub(r"<[^>]+>", " ", raw))
        elif ext == ".rtf":
            text = re.sub(r"\\'[0-9a-fA-F]{2}|\\[a-z]+-?\d* ?|[{}]", "", raw)
        elif ext in (".yaml", ".yml"):
            try:
                import yaml  # type: ignore
                return _json_blocks(yaml.safe_load(raw))
            except ImportError:
                text = raw
        else:
            text = raw
        return _line_blocks(text)


class CsvParser(Parser):
    extensions = (".csv",)

    def parse(self, path: str, name: str) -> list[dict[str, Any]]:
        with open(path, newline="", encoding="utf-8-sig", errors="replace") as f:
            rows = list(csv.reader(f))
        if not rows:
            return []
        headers = rows[0]
        result = [block(" | ".join(headers), {"row": 1}, "Columns", "table_header")]
        for i, row in enumerate(rows[1:], 2):
            values = [f"{headers[j] if j < len(headers) else 'Column '+str(j+1)}: {v}" for j, v in enumerate(row) if v.strip()]
            if values:
                result.append(block(" | ".join(values), {"row": i}, "Data", "table_row"))
        return result


class PdfParser(Parser):
    extensions = (".pdf",)

    def parse(self, path: str, name: str) -> list[dict[str, Any]]:
        try:
            from pypdf import PdfReader  # type: ignore
        except ImportError as exc:
            raise RuntimeError("PDF support requires pypdf. Install requirements.txt.") from exc
        reader = PdfReader(path, strict=False)
        if reader.is_encrypted:
            try:
                if not reader.decrypt(""):
                    raise ValueError("This PDF is password-protected. Remove its password and upload it again.")
            except Exception as exc:
                raise ValueError("This PDF is password-protected. Remove its password and upload it again.") from exc
        out = []
        for n, page in enumerate(reader.pages, 1):
            text = page.extract_text() or ""
            # A few stray glyphs on an otherwise scanned page are not meaningful text.
            if len(re.sub(r"\s", "", text)) >= 20:
                out.extend(_line_blocks(text, {"page": n}, method="embedded_text", source_type="text"))
                continue
            try:
                ocr_text = _ocr_pdf_page(path, n - 1)
            except RuntimeError:
                raise
            except Exception as exc:
                logger.exception("PDF OCR failed for %s page %s", name, n)
                raise RuntimeError(f"OCR failed on page {n} of {name}: {exc}") from exc
            if ocr_text.strip():
                out.extend(_line_blocks(ocr_text, {"page": n}, method="ocr", confidence=0.65, source_type="ocr"))
            elif text.strip():
                # Preserve sparse embedded content when OCR cannot improve it.
                out.extend(_line_blocks(text, {"page": n}, method="embedded_text", source_type="text"))
            else:
                out.append(block("OCR found no readable text on this page.", {"page": n}, "", "ocr_required", confidence=0.0, extraction_method="ocr", source_type="ocr"))
        return out


def _ocr_pdf_page(path: str, page_index: int) -> str:
    """Render one scanned PDF page and OCR it without requiring a separate poppler install."""
    try:
        import fitz  # type: ignore
        from PIL import Image
        import pytesseract  # type: ignore
    except ImportError as exc:
        raise RuntimeError("Scanned PDF OCR requires PyMuPDF, Pillow, pytesseract, and the Tesseract system executable. Install Python packages from requirements.txt and install Tesseract.") from exc
    try:
        with fitz.open(path) as doc:
            page = doc.load_page(page_index)
            pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
            image = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)
            try:
                return pytesseract.image_to_string(image)
            except pytesseract.TesseractNotFoundError as exc:
                raise RuntimeError("Tesseract was not found. Install the Tesseract system executable and retry OCR.") from exc
    except Exception as exc:
        raise RuntimeError(f"Unable to render/OCR PDF page {page_index + 1}: {exc}") from exc


class DocxParser(Parser):
    extensions = (".docx",)

    def parse(self, path: str, name: str) -> list[dict[str, Any]]:
        try:
            from docx import Document as WordDocument  # type: ignore
        except ImportError as exc:
            raise RuntimeError("DOCX support requires python-docx. Install requirements.txt.") from exc
        doc = WordDocument(path)
        out = []
        section = ""
        para_no = 0
        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue
            para_no += 1
            if p.style and p.style.name.startswith("Heading"):
                section = text
            out.append(block(text, {"paragraph": para_no}, section, "heading" if p.style and p.style.name.startswith("Heading") else "text"))
        for ti, table in enumerate(doc.tables, 1):
            for ri, row in enumerate(table.rows, 1):
                vals = [cell.text.strip() for cell in row.cells]
                out.append(block(" | ".join(vals), {"table": ti, "row": ri}, section, "table_row"))
        return out


class PptxParser(Parser):
    extensions = (".pptx",)

    def parse(self, path: str, name: str) -> list[dict[str, Any]]:
        try:
            from pptx import Presentation  # type: ignore
        except ImportError as exc:
            raise RuntimeError("PPTX support requires python-pptx. Install requirements.txt.") from exc
        out = []
        for sn, slide in enumerate(Presentation(path).slides, 1):
            for oi, shape in enumerate(slide.shapes, 1):
                if getattr(shape, "has_text_frame", False) and shape.text.strip():
                    out.append(block(shape.text.strip(), {"slide": sn, "object": oi}, "", "text"))
                if getattr(shape, "has_table", False):
                    for ri, row in enumerate(shape.table.rows, 1):
                        out.append(block(" | ".join(c.text for c in row.cells), {"slide": sn, "object": oi, "row": ri}, "", "table_row"))
            try:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    out.append(block(notes, {"slide": sn, "object": "speaker notes"}, "", "speaker_notes"))
            except (AttributeError, KeyError):
                pass
        return out


class XlsxParser(Parser):
    extensions = (".xlsx",)

    def parse(self, path: str, name: str) -> list[dict[str, Any]]:
        try:
            from openpyxl import load_workbook  # type: ignore
        except ImportError as exc:
            raise RuntimeError("XLSX support requires openpyxl. Install requirements.txt.") from exc
        wb = load_workbook(path, data_only=False, read_only=True)
        cached = load_workbook(path, data_only=True, read_only=True)
        out = []
        for ws in wb.worksheets:
            values = cached[ws.title]
            for row in ws.iter_rows():
                parts = []
                for cell in row:
                    if cell.value is None:
                        continue
                    v = str(cell.value)
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        calc = values[cell.coordinate].value
                        v += f" (calculated value: {calc})" if calc is not None else " (formula; calculated value unavailable)"
                    parts.append(f"{cell.coordinate}: {v}")
                if parts:
                    out.append(block(" | ".join(parts), {"sheet": ws.title, "range": f"{row[0].coordinate}:{row[-1].coordinate}"}, ws.title, "table_row"))
        return out


class ImageParser(Parser):
    extensions = (".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff", ".bmp")

    def parse(self, path: str, name: str) -> list[dict[str, Any]]:
        try:
            from PIL import Image  # type: ignore
            import pytesseract  # type: ignore
            image = Image.open(path)
            text = pytesseract.image_to_string(image)
        except ImportError as exc:
            raise RuntimeError("Image OCR requires Pillow, pytesseract, and the Tesseract system package.") from exc
        if not text.strip():
            return [block("No readable text detected. Visual interpretation is not configured.", {"image": name}, "", "image", confidence=0.0)]
        return _line_blocks(text, {"image": name}, method="ocr", confidence=0.65, source_type="ocr")


def _line_blocks(text: str, base: dict[str, Any] | None = None, method: str = "parser", confidence: float = 1.0, source_type: str = "text") -> list[dict[str, Any]]:
    out = []
    for i, line in enumerate(text.splitlines(), 1):
        line = re.sub(r"\s+", " ", line).strip()
        if line:
            loc = dict(base or {})
            loc.setdefault("line", i)
            out.append(block(line, loc, "", "ocr_text" if source_type == "ocr" else "text", extraction_method=method, confidence=confidence, source_type=source_type))
    return out


def _json_blocks(value: Any, path: str = "$", depth: int = 0) -> list[dict[str, Any]]:
    if isinstance(value, dict):
        out = []
        for key, child in value.items():
            out.extend(_json_blocks(child, f"{path}.{key}", depth + 1))
        return out or [block("{}", {"json_path": path}, path, "structured_value")]
    if isinstance(value, list):
        out = []
        for i, child in enumerate(value):
            out.extend(_json_blocks(child, f"{path}[{i}]", depth + 1))
        return out
    return [block(f"{path.rsplit('.', 1)[-1].split('[')[0]}: {value}", {"json_path": path}, path, "structured_value")]


def _xml_blocks(root: ET.Element) -> list[dict[str, Any]]:
    out = []
    for el in root.iter():
        path = "/" + "/".join(e.tag for e in list(root.iter()) if e is el)
        attrs = " ".join(f"{k}={v}" for k, v in el.attrib.items())
        val = (el.text or "").strip()
        content = " ".join(x for x in (attrs, val) if x)
        if content:
            out.append(block(f"{el.tag}: {content}", {"xpath": f"//{el.tag}"}, el.tag, "structured_value"))
    return out


PARSERS: dict[str, Parser] = {}
for _parser in (TextParser(), CsvParser(), PdfParser(), DocxParser(), PptxParser(), XlsxParser(), ImageParser()):
    for _ext in _parser.extensions:
        PARSERS[_ext] = _parser


def parse_file(path: str, name: str) -> list[dict[str, Any]]:
    ext = Path(name).suffix.lower()
    parser = PARSERS.get(ext)
    if parser is None:
        raise ValueError(f"Unsupported file type: {ext or 'unknown'}")
    return parser.parse(path, name)
